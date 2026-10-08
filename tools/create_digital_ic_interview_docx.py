from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "interview" / "数字IC校招项目面试模拟与参考回答.md"
OUTPUT = ROOT / "数字IC校招项目面试模拟与参考回答.docx"

FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"
BLACK = RGBColor(0, 0, 0)
NAVY = RGBColor(31, 78, 121)
MID_BLUE = RGBColor(47, 84, 150)
GRAY = RGBColor(89, 89, 89)
LIGHT_BLUE = "EAF2F8"


def set_font(run, size, color=BLACK, bold=False, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def set_style_font(style, size, color=BLACK, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style._element.get_or_add_rPr()
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def set_spacing(paragraph, before=0, after=0, line=1.18):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def set_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def clean_inline(text):
    return text.replace("`", "").replace("**", "").strip()


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(paragraph, before=0, after=0, line=1.0)
    set_font(paragraph.add_run("第 "), 9, GRAY)

    begin_run = paragraph.add_run()
    set_font(begin_run, 9, GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    set_font(instr_run, 9, GRAY)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)

    separate_run = paragraph.add_run()
    set_font(separate_run, 9, GRAY)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    display_run = paragraph.add_run("1")
    set_font(display_run, 9, GRAY)

    end_run = paragraph.add_run()
    set_font(end_run, 9, GRAY)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)

    set_font(paragraph.add_run(" 页"), 9, GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    set_style_font(styles["Normal"], 10.5, BLACK)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.18

    set_style_font(styles["Title"], 20, BLACK, True)
    styles["Title"].paragraph_format.space_after = Pt(9)
    styles["Title"].paragraph_format.line_spacing = 1.0

    set_style_font(styles["Heading 1"], 15, NAVY, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].paragraph_format.keep_with_next = True

    set_style_font(styles["Heading 2"], 12.5, NAVY, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.keep_with_next = True

    set_style_font(styles["Heading 3"], 11, MID_BLUE, True)
    styles["Heading 3"].paragraph_format.space_before = Pt(6)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)
    styles["Heading 3"].paragraph_format.keep_with_next = True

    set_style_font(styles["List Bullet"], 10.3, BLACK)
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.24)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.14)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    styles["List Bullet"].paragraph_format.line_spacing = 1.15

    if "Interview Label" not in [style.name for style in styles]:
        label_style = styles.add_style("Interview Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        label_style = styles["Interview Label"]
    set_style_font(label_style, 10.5, NAVY, True)
    label_style.paragraph_format.space_before = Pt(3)
    label_style.paragraph_format.space_after = Pt(2)
    label_style.paragraph_format.keep_with_next = True

    if "Contents Entry" not in [style.name for style in styles]:
        toc_style = styles.add_style("Contents Entry", WD_STYLE_TYPE.PARAGRAPH)
    else:
        toc_style = styles["Contents Entry"]
    set_style_font(toc_style, 11, BLACK)
    toc_style.paragraph_format.left_indent = Inches(0.18)
    toc_style.paragraph_format.space_after = Pt(7)

    add_page_field(section.footer.paragraphs[0])


def add_cover(doc):
    spacer = doc.add_paragraph()
    set_spacing(spacer, after=60, line=1.0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("数字 IC 校招项目面试模拟与参考回答"), 20, BLACK, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(subtitle, after=16, line=1.0)
    set_font(subtitle.add_run("RISC V Audio DSP SoC RTL 设计与验证"), 12, NAVY, True)

    audience = doc.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(audience, after=22, line=1.0)
    set_font(audience.add_run("面向数字 IC 前端校招  24 道主问题  53 道递进追问"), 10.5, GRAY)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    summary.paragraph_format.left_indent = Inches(0.72)
    summary.paragraph_format.right_indent = Inches(0.72)
    set_spacing(summary, before=4, after=0, line=1.3)
    set_font(
        summary.add_run(
            "本材料围绕当前仓库中已经实现并通过仿真验证的 CPU、APB、BRAM、DMA、CDC、FIR 与 I2S 展开。"
            "回答区分 RTL 仿真事实与尚待完成的综合、时序和板级验证。"
        ),
        11,
        BLACK,
    )
    doc.add_page_break()


def add_contents(doc):
    heading = doc.add_paragraph("目录", style="Heading 1")
    heading.paragraph_format.page_break_before = False
    entries = [
        "使用说明与项目介绍",
        "第一轮  项目全局与个人工作",
        "第二轮  总线外设与存储系统",
        "第三轮  CPU CSR Trap 与中断",
        "第四轮  DMA 仲裁与数据搬运",
        "第五轮  CDC 与异步 FIFO",
        "第六轮  FIR 定点音频与 I2S",
        "第七轮  验证调试与项目复盘",
        "面试收尾准备",
    ]
    for index, entry in enumerate(entries, 1):
        paragraph = doc.add_paragraph(style="Contents Entry")
        set_spacing(paragraph, after=7, line=1.1)
        set_font(paragraph.add_run(f"{index:02d}  {entry}"), 11, BLACK)
    note = doc.add_paragraph()
    set_spacing(note, before=14, after=0, line=1.2)
    set_font(note.add_run("建议先遮住参考回答进行口述，再按回答要点复盘。"), 10, GRAY)
    doc.add_page_break()


def add_heading(doc, level, text):
    paragraph = doc.add_paragraph(clean_inline(text), style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        paragraph.paragraph_format.page_break_before = True
    if level == 2 and text.startswith("Q"):
        set_shading(paragraph, LIGHT_BLUE)
        paragraph.paragraph_format.left_indent = Inches(0.04)
        paragraph.paragraph_format.right_indent = Inches(0.04)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(paragraph, after=5, line=1.18)
    set_font(paragraph.add_run(clean_inline(text)), 10.5, BLACK)
    return paragraph


def add_label(doc, text):
    paragraph = doc.add_paragraph(style="Interview Label")
    set_font(paragraph.add_run(clean_inline(text)), 10.5, NAVY, True)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(paragraph.add_run(clean_inline(text)), 10.3, BLACK)
    return paragraph


def append_markdown(doc, source_text):
    lines = source_text.splitlines()
    title_skipped = False
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not title_skipped:
            title_skipped = True
            continue
        if line.startswith("### "):
            add_heading(doc, 3, line[4:])
        elif line.startswith("## "):
            add_heading(doc, 2, line[3:])
        elif line.startswith("# "):
            add_heading(doc, 1, line[2:])
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif line.startswith("**") and line.endswith("**"):
            add_label(doc, line)
        else:
            add_body(doc, line)


def main():
    source_text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)
    append_markdown(doc, source_text)

    doc.core_properties.title = "数字 IC 校招项目面试模拟与参考回答"
    doc.core_properties.subject = "RISC V Audio DSP SoC RTL 设计与验证"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "数字 IC 校招 SystemVerilog RISC V APB DMA CDC FIR I2S"
    doc.settings.element.set(qn("w:updateFields"), "true")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
