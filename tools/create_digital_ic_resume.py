from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "数字IC前端项目经历.docx"
FONT = "Microsoft YaHei"
NAVY = RGBColor(31, 78, 121)
GRAY = RGBColor(89, 89, 89)


def set_run_font(run, size, color=None, bold=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=5, after=2, line=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, 11, NAVY, bold=True)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    set_paragraph_spacing(paragraph, after=2, line=1.1)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5)
    return paragraph


def add_compact_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    set_paragraph_spacing(paragraph, after=1, line=1.05)
    run = paragraph.add_run(text)
    set_run_font(run, 9.7)
    return paragraph


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, after=1, line=1.0)
    set_run_font(title.add_run("数字 IC 前端项目经历"), 18, NAVY, bold=True)

    project = doc.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(project, after=6, line=1.0)
    set_run_font(project.add_run("RISC-V Audio DSP SoC RTL 设计与验证"), 12, GRAY, bold=True)

    stack = doc.add_table(rows=1, cols=1)
    stack.autofit = True
    stack_cell = stack.cell(0, 0)
    set_cell_shading(stack_cell, "EAF2F8")
    p = stack_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=2, line=1.0)
    lead = p.add_run("技术栈：")
    set_run_font(lead, 9.7, NAVY, bold=True)
    details = p.add_run("SystemVerilog | RISC-V RV32IM | APB | DMA | BRAM | CDC/Async FIFO | Q1.15 FIR | Icarus Verilog")
    set_run_font(details, 9.7, GRAY)

    add_section_heading(doc, "项目简介")
    overview = doc.add_paragraph()
    set_paragraph_spacing(overview, after=2, line=1.12)
    set_run_font(
        overview.add_run(
            "设计并验证一套面向 FPGA 原型的 RISC-V Audio DSP SoC RTL，覆盖 CPU、片上存储与 APB 外设、DMA、多时钟域音频数据通路和定点 FIR；当前工作以 RTL 仿真验证为边界。"
        ),
        10.5,
    )

    add_section_heading(doc, "核心工作与成果")
    bullets = [
        "设计 APB 外设子系统及地址译码，完成 GPIO、Timer、UART TX/RX、寄存器接口和 16 深度收发 FIFO；支持非法访问 PSLVERR、UART 错误状态及可屏蔽中断。",
        "实现 64 KiB 片上 BRAM、单 outstanding CPU 请求/响应总线和 CPU-to-APB Bridge；支持 32 bit 读写、字节写使能、APB Setup/Access 时序与错误透传。",
        "自研模块化多周期 RV32IM 内核，覆盖取指、译码、执行、访存、写回，支持 Machine CSR、Trap、MUL/DIV/REM 及 Timer、DMA、UART 中断汇聚。",
        "实现单通道 DMA 及 CPU/DMA BRAM 仲裁，基于读写状态机完成 BRAM 到 BRAM 和音频 FIFO 端点搬运；覆盖零长度、非对齐、忙态重复启动及 RW1C 状态清除。",
        "构建 50 MHz 系统域至 12.5 MHz 音频域的 CDC 数据通路：参数化 Gray-pointer Async FIFO 与 16 阶 Q1.15 FIR；完成 BRAM → DMA → FIFO → FIR → FIFO → DMA → BRAM 的 40 样本端到端逐样本黄金模型比对。",
    ]
    for text in bullets:
        add_bullet(doc, text)

    verification = doc.add_paragraph()
    set_paragraph_spacing(verification, before=2, after=1, line=1.05)
    set_run_font(verification.add_run("验证："), 10, NAVY, bold=True)
    set_run_font(
        verification.add_run("基于 Icarus Verilog 完成 33 个 SystemVerilog Testbench 回归，全部 PASS；apb_subsystem 与 soc_top 静态 elaboration 通过。"),
        10,
    )

    add_section_heading(doc, "精简替换版")
    compact_bullets = [
        "使用 SystemVerilog 完成含多周期 RV32IM、64 KiB BRAM、APB 外设和 DMA 的 SoC RTL 设计，支持 CSR/Trap、异常处理及 Timer/DMA/UART 中断。",
        "设计 Gray-pointer Async FIFO 和 16 阶 Q1.15 FIR，完成 50 MHz/12.5 MHz 双时钟域 BRAM-DMA-FIFO-FIR-DMA 音频链路，40 样本逐样本与黄金模型一致。",
        "搭建单元至 SoC 的 SystemVerilog 自检回归，覆盖协议时序、异常状态、FIFO 满空/回绕/独立复位及端到端数据搬运；33 个 Testbench 全部 PASS。",
    ]
    for text in compact_bullets:
        add_compact_bullet(doc, text)

    doc.core_properties.title = "数字 IC 前端项目经历"
    doc.core_properties.subject = "RISC-V Audio DSP SoC RTL 设计与验证"
    doc.core_properties.author = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
